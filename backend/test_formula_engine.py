from pathlib import Path
from decimal import Decimal

import pytest
from docx import Document

from formula_engine import FormulaError, evaluate_template_formulas, select_fields
from generation import extract_placeholders_from_docx, render_document


def make_template(path: Path, formulas: list[str]) -> Path:
    document = Document()
    for formula in formulas:
        document.add_paragraph(formula)
    document.add_paragraph("Result: {{total_fixed_monthly}}")
    document.save(path)
    return path


def test_dependency_resolution_and_pf_boundaries(tmp_path):
    template = make_template(tmp_path / "rules.docx", [
        "[CALC(basic_monthly = {{ctc_total}} / 24)]",
        "[CALC(pf_monthly = MIN({{basic_monthly}}, 25000) * 0.12)]",
        "[CALC(hra_monthly = {{basic_monthly}} * 0.10)]",
        "[CALC(total_fixed_monthly = {{basic_monthly}} + {{hra_monthly}} + {{pf_monthly}})]",
    ])
    # Basic values: 20k, 25k, 25,001, and 30k. PF cap is defined by template.
    for basic, expected_pf in [(20000, 2400), (25000, 3000), (25001, 3000), (30000, 3000)]:
        values, _ = evaluate_template_formulas(template, {"ctc_total": basic * 24})
        assert values["basic_monthly"] == basic
        assert values["pf_monthly"] == expected_pf


def test_if_min_max_and_inline_formula(tmp_path):
    template = make_template(tmp_path / "if.docx", [
        "[CALC(basic_monthly = {{ctc_total}} / 24)]",
        "[CALC(pf_monthly = IF({{basic_monthly}} <= 25000, {{basic_monthly}} * 0.12, 25000 * 0.12))]",
        "Inline PF: [CALC(MAX({{pf_monthly}}, 1800))]",
    ])
    values, replacements = evaluate_template_formulas(template, {"ctc_total": 600000})
    assert values["basic_monthly"] == 25000
    assert values["pf_monthly"] == 3000
    assert "3000" in replacements.values()


@pytest.mark.parametrize(("expression", "expected"), [
    ("ROUND(10.4, 0)", Decimal("10")),
    ("ROUND(10.5, 0)", Decimal("11")),
    ("ROUND(10.6, 0)", Decimal("11")),
    ("ROUND(123.456, 2)", Decimal("123.46")),
    ("ROUND(123.454, 2)", Decimal("123.45")),
    ("ROUND(123.455, 2)", Decimal("123.46")),
    ("ROUND(-10.5, 0)", Decimal("-11")),
    ("ROUND(MIN(30000, 25000) * 0.12, 0)", Decimal("3000")),
])
def test_round_uses_excel_half_away_from_zero(tmp_path, expression, expected):
    template = make_template(tmp_path / "round.docx", [f"[CALC(result = {expression})]"])
    values, _ = evaluate_template_formulas(template, {})
    assert values["result"] == expected


def test_round_resolves_placeholders_and_preserves_existing_min(tmp_path):
    template = make_template(tmp_path / "round_placeholder.docx", [
        "[CALC(basic_monthly = ROUND({{ctc_total}} / 24, 0))]",
        "[CALC(pf_monthly = MIN({{basic_monthly}}, 25000) * 0.12)]",
    ])
    # 240,012 / 24 = 10,000.5, which must round up to 10,001.
    values, replacements = evaluate_template_formulas(template, {"ctc_total": 240012})
    assert values["basic_monthly"] == Decimal("10001")
    assert values["pf_monthly"] == Decimal("1200.12")
    assert "10001" in replacements.values()


def test_round_rejects_non_integer_digits(tmp_path):
    template = make_template(tmp_path / "invalid_round.docx", ["[CALC(result = ROUND(10.5, 0.5))]"])
    with pytest.raises(FormulaError, match="ROUND digits must be an integer"):
        evaluate_template_formulas(template, {})


def test_undefined_variable_is_clear(tmp_path):
    template = make_template(tmp_path / "undefined.docx", ["[CALC(result = {{unknown_variable}} + 100)]"])
    with pytest.raises(FormulaError, match="undefined variable: unknown_variable"):
        evaluate_template_formulas(template, {"ctc_total": 100})


def test_invalid_formula_is_clear(tmp_path):
    template = make_template(tmp_path / "invalid.docx", ["[CALC(result = {{ctc_total}} /)]"])
    with pytest.raises(FormulaError, match="Invalid formula"):
        evaluate_template_formulas(template, {"ctc_total": 100})


def test_circular_dependency_is_clear(tmp_path):
    template = make_template(tmp_path / "cycle.docx", [
        "[CALC(a = {{b}} + 1)]",
        "[CALC(b = {{a}} + 1)]",
    ])
    with pytest.raises(FormulaError, match="Circular calculation dependency detected"):
        evaluate_template_formulas(template, {})


def test_docx_generation_preserves_template_and_reacts_to_changed_formula(tmp_path):
    template = make_template(tmp_path / "integration.docx", [
        "[CALC(total_fixed_monthly = {{ctc_total}} / 24)]",
    ])
    output_one, _ = render_document(template, {"ctc_total": 240000}, "formula_one", "docx")
    assert output_one.exists()
    # Changing only the DOCX formula changes the result; backend code is untouched.
    make_template(template, ["[CALC(total_fixed_monthly = {{ctc_total}} / 12)]"])
    output_two, _ = render_document(template, {"ctc_total": 240000}, "formula_two", "docx")
    assert output_two.exists()
    assert output_one.read_bytes() != output_two.read_bytes()


def test_currency_inputs_and_calculations_use_indian_comma_grouping(tmp_path):
    template = make_template(tmp_path / "currency_format.docx", [
        "Annual CTC: {{ctc_total}}",
        "[CALC(basic_monthly = {{ctc_total}} / 24)]",
        "[CALC(gross_monthly = {{ctc_total}} / 12)]",
        "Annual from inline formula: [CALC({{basic_monthly}} * 12)]",
    ])
    output, _ = render_document(
        template,
        {"ctc_total": 240000},
        "currency_format",
        "docx",
        currency_fields={"ctc_total", "basic_monthly", "gross_monthly"},
    )
    generated_text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "2,40,000" in generated_text
    assert "10,000" in generated_text
    assert "20,000" in generated_text
    assert "1,20,000" in generated_text


def test_docx_select_tag_becomes_a_dropdown_field_and_renders_selection(tmp_path):
    template = make_template(tmp_path / "dropdown.docx", [
        "Annual medical insurance: [SELECT(insurance_annual, 8000, 15000)]",
    ])
    fields = select_fields(template)
    assert fields == {"insurance_annual": ("8000", "15000")}
    placeholders = extract_placeholders_from_docx(template)
    insurance = next(field for field in placeholders if field["name"] == "insurance_annual")
    assert insurance["options"] == ["8000", "15000"]

    output, _ = render_document(
        template,
        {"insurance_annual": 15000},
        "dropdown_output",
        "docx",
        currency_fields={"insurance_annual"},
    )
    generated = Document(output)
    assert "[SELECT(" not in "\n".join(paragraph.text for paragraph in generated.paragraphs)
    assert "15,000" in "\n".join(paragraph.text for paragraph in generated.paragraphs)


def test_select_numeric_and_expression_options_feed_downstream_calculations(tmp_path):
    template = make_template(tmp_path / "pf_select.docx", [
        "[CALC(basic_monthly = {{ctc_total}} / 24)]",
        "[SELECT(pf_monthly, 3000, {{basic_monthly}} * 0.12)]",
        "[CALC(pf_annual = {{pf_monthly}} * 12)]",
    ])

    fixed, _ = evaluate_template_formulas(
        template, {"ctc_total": 720000}, {"pf_monthly": "3000"}
    )
    assert fixed["basic_monthly"] == Decimal("30000")
    assert fixed["pf_monthly"] == Decimal("3000")
    assert fixed["pf_annual"] == Decimal("36000")

    percentage, _ = evaluate_template_formulas(
        template,
        {"ctc_total": 720000},
        {"pf_monthly": "{{basic_monthly}} * 0.12"},
    )
    assert percentage["pf_monthly"] == Decimal("3600")
    assert percentage["pf_annual"] == Decimal("43200")


def test_select_nested_calc_option_feeds_downstream_calculations(tmp_path):
    template = make_template(tmp_path / "pf_select_nested_calc.docx", [
        "[CALC(basic_monthly = {{ctc_total}} / 24)]",
        "[SELECT(pf_monthly, 3000, CALC[0.12 * {{basic_monthly}}])]",
        "[CALC(pf_annual = {{pf_monthly}} * 12)]",
    ])
    assert select_fields(template) == {
        "pf_monthly": ("3000", "CALC[0.12 * {{basic_monthly}}]")
    }
    fixed, _ = evaluate_template_formulas(
        template, {"ctc_total": 720000}, {"pf_monthly": "3000"}
    )
    assert fixed["pf_monthly"] == Decimal("3000")
    assert fixed["pf_annual"] == Decimal("36000")

    calculated, _ = evaluate_template_formulas(
        template,
        {"ctc_total": 720000},
        {"pf_monthly": "CALC[0.12 * {{basic_monthly}}]"},
    )
    assert isinstance(calculated["pf_monthly"], Decimal)
    assert calculated["pf_monthly"] == Decimal("3600")
    assert calculated["pf_annual"] == Decimal("43200")


def test_select_expression_recalculates_when_its_dependency_changes(tmp_path):
    template = make_template(tmp_path / "pf_select_recalculate.docx", [
        "[CALC(basic_monthly = {{ctc_total}} / 24)]",
        "[SELECT(pf_monthly, 3000, {{basic_monthly}} * 0.12)]",
        "[CALC(pf_annual = {{pf_monthly}} * 12)]",
    ])
    values, _ = evaluate_template_formulas(
        template,
        {"ctc_total": 480000},
        {"pf_monthly": "{{basic_monthly}} * 0.12"},
    )
    assert values["basic_monthly"] == Decimal("20000")
    assert values["pf_monthly"] == Decimal("2400")
    assert values["pf_annual"] == Decimal("28800")


def test_nested_select_calc_recalculates_when_its_dependency_changes(tmp_path):
    template = make_template(tmp_path / "pf_select_nested_recalculate.docx", [
        "[CALC(basic_monthly = {{ctc_total}} / 24)]",
        "[SELECT(pf_monthly, 3000, CALC[0.12 * {{basic_monthly}}])]",
        "[CALC(pf_annual = {{pf_monthly}} * 12)]",
    ])
    values, _ = evaluate_template_formulas(
        template,
        {"ctc_total": 480000},
        {"pf_monthly": "CALC[0.12 * {{basic_monthly}}]"},
    )
    assert values["pf_monthly"] == Decimal("2400")
    assert values["pf_annual"] == Decimal("28800")


def test_invalid_nested_select_calc_has_clear_error(tmp_path):
    template = make_template(tmp_path / "invalid_nested_pf_select.docx", [
        "[SELECT(pf_monthly, 3000, CALC[0.12 * {{unknown_basic}}])]",
        "[CALC(pf_annual = {{pf_monthly}} * 12)]",
    ])
    with pytest.raises(FormulaError, match="Dropdown 'pf_monthly' selected expression.*unknown_basic"):
        evaluate_template_formulas(
            template,
            {},
            {"pf_monthly": "CALC[0.12 * {{unknown_basic}}]"},
        )


def test_select_expression_is_numeric_even_without_a_downstream_formula(tmp_path):
    template = make_template(tmp_path / "pf_select_direct.docx", [
        "[CALC(basic_monthly = {{ctc_total}} / 24)]",
        "[SELECT(pf_monthly, 3000, {{basic_monthly}} * 0.12)]",
    ])
    values, _ = evaluate_template_formulas(
        template,
        {"ctc_total": 720000},
        {"pf_monthly": "{{basic_monthly}} * 0.12"},
    )
    assert values["pf_monthly"] == Decimal("3600")


def test_pf_select_value_is_rendered_with_indian_currency_commas(tmp_path):
    template = make_template(tmp_path / "pf_select_currency.docx", [
        "PF: [SELECT(pf_monthly, 3000, {{basic_monthly}} * 0.12)]",
        "[CALC(basic_monthly = {{ctc_total}} / 24)]",
    ])
    output, _ = render_document(
        template,
        {"ctc_total": 720000, "pf_monthly": "3000"},
        "pf_select_currency",
        "docx",
        currency_fields={"pf_monthly", "basic_monthly"},
    )
    generated_text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "PF: 3,000" in generated_text


def test_nested_calc_select_value_is_evaluated_before_currency_formatting(tmp_path):
    template = make_template(tmp_path / "pf_nested_select_currency.docx", [
        "PF: [SELECT(pf_monthly, 3000, CALC[0.12 * {{basic_monthly}}])]",
        "[CALC(basic_monthly = {{ctc_total}} / 24)]",
    ])
    output, _ = render_document(
        template,
        {"ctc_total": 720000, "pf_monthly": "CALC[0.12 * {{basic_monthly}}]"},
        "pf_nested_select_currency",
        "docx",
        currency_fields={"pf_monthly", "basic_monthly"},
    )
    generated_text = "\n".join(paragraph.text for paragraph in Document(output).paragraphs)
    assert "PF: 3,600" in generated_text


def test_invalid_select_expression_has_clear_error(tmp_path):
    template = make_template(tmp_path / "invalid_pf_select.docx", [
        "[SELECT(pf_monthly, 3000, {{unknown_basic}} * 0.12)]",
        "[CALC(pf_annual = {{pf_monthly}} * 12)]",
    ])
    with pytest.raises(FormulaError, match="Dropdown 'pf_monthly' selected expression.*unknown_basic"):
        evaluate_template_formulas(
            template,
            {},
            {"pf_monthly": "{{unknown_basic}} * 0.12"},
        )


def test_placeholders_follow_first_appearance_in_document_order(tmp_path):
    template = make_template(tmp_path / "sequence.docx", [
        "Date: {{offer_date}}",
        "Employee: {{full_name}}",
        "CTC: {{ctc_total}}",
        "Designation: {{designation}}",
        "[CALC(basic_monthly = {{ctc_total}} / 24)]",
    ])
    placeholders = extract_placeholders_from_docx(template)
    input_names = [field["name"] for field in placeholders if not field["calculated"]]
    assert input_names == ["offer_date", "full_name", "ctc_total", "designation", "total_fixed_monthly"]
    assert next(field for field in placeholders if field["name"] == "basic_monthly")["calculated"] is True
