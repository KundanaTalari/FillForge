import os
import uuid
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from db import init_db, insert_document, get_all_documents
from storage import init_storage, get_template_path
from generation import extract_placeholders_from_docx

def create_offer_letter_template(dest_path: Path):
    doc = Document()

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("FILLFORGE TECHNOLOGIES PVT. LTD.")
    run_title.bold = True
    run_title.font.size = Pt(18)
    run_title.font.color.rgb = RGBColor(30, 41, 59)

    sub = doc.add_paragraph()
    sub.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub.add_run("LETTER OF EMPLOYMENT OFFER & COMPENSATION BREAKDOWN")
    run_sub.bold = True
    run_sub.font.size = Pt(13)
    run_sub.font.color.rgb = RGBColor(79, 70, 229)

    doc.add_paragraph().paragraph_format.space_after = Pt(12)

    # Header details
    p_date = doc.add_paragraph()
    p_date.add_run("Date: ").bold = True
    p_date.add_run("{{date_of_joining}}")

    p_to = doc.add_paragraph()
    p_to.add_run("To,\n").bold = True
    p_to.add_run("{{first_name}} {{last_name}}\n")
    p_to.add_run("Employee ID: {{employee_id}}\n")
    p_to.add_run("Email: {{email}}\n")
    p_to.add_run("Contact: {{phone}}")

    doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Greeting & Offer
    p_greet = doc.add_paragraph()
    p_greet.add_run("Dear {{first_name}},\n\n")
    p_greet.add_run(
        "We are excited to extend an offer of employment for the position of {{designation}} "
        "in the {{department}} Department at FillForge Technologies. Your start date will be "
        "{{date_of_joining}} at {{joining_time}}."
    )

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Compensation Section
    p_comp = doc.add_paragraph()
    run_comp = p_comp.add_run("Annexure A: Compensation & Benefits Structure")
    run_comp.bold = True
    run_comp.font.size = Pt(12)

    # Table
    table = doc.add_table(rows=1, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Salary Component"
    hdr_cells[1].text = "Annual (₹)"
    hdr_cells[2].text = "Monthly (₹)"

    for cell in hdr_cells:
        for p in cell.paragraphs:
            for r in p.runs:
                r.bold = True

    components = [
        ("Basic Salary", "{{annual_basic}}", "{{basic_per_month}}"),
        ("House Rent Allowance (HRA)", "{{annual_hra}}", "{{hra_per_month}}"),
        ("Provident Fund (PF)", "{{pf_per_year}}", "{{pf_per_month}}"),
        ("Gratuity (4.81%)", "{{gratuity_per_year}}", "{{gratuity_per_month}}"),
        ("Special Allowance", "{{special_allowance}}", "{{monthly_special_allowance}}"),
        ("Total CTC (Cost to Company)", "{{ctc_total}}", "[CALC({{ctc_total}}/12)]"),
    ]

    for comp, annual, monthly in components:
        row_cells = table.add_row().cells
        row_cells[0].text = comp
        row_cells[1].text = annual
        row_cells[2].text = monthly

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Closing
    p_close = doc.add_paragraph()
    p_close.add_run(
        "Please sign and return the duplicate copy of this letter as an acceptance of our offer.\n\n"
        "Sincerely,\n"
        "Human Resources Department\n"
        "FillForge Technologies Pvt. Ltd."
    )

    doc.save(str(dest_path))


def create_internship_certificate_template(dest_path: Path):
    doc = Document()

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title.add_run("CERTIFICATE OF INTERNSHIP")
    run_title.bold = True
    run_title.font.size = Pt(20)
    run_title.font.color.rgb = RGBColor(30, 41, 59)

    doc.add_paragraph().paragraph_format.space_after = Pt(20)

    p = doc.add_paragraph()
    p.paragraph_format.line_spacing = 1.6
    p.add_run("This is to certify that ").font.size = Pt(12)
    r_name = p.add_run("{{first_name}} {{last_name}}")
    r_name.bold = True
    r_name.font.size = Pt(13)
    p.add_run(" (Intern ID: {{employee_id}}) has successfully served as an intern in the ").font.size = Pt(12)
    r_dept = p.add_run("{{department}}")
    r_dept.bold = True
    r_dept.font.size = Pt(12)
    p.add_run(" department from ").font.size = Pt(12)
    r_date = p.add_run("{{date_of_joining}}")
    r_date.bold = True
    r_date.font.size = Pt(12)
    p.add_run(" under the designation of ").font.size = Pt(12)
    r_desig = p.add_run("{{designation}}")
    r_desig.bold = True
    r_desig.font.size = Pt(12)
    p.add_run(" with a monthly stipend of ₹{{salary}}.").font.size = Pt(12)

    doc.add_paragraph().paragraph_format.space_after = Pt(30)

    p_sig = doc.add_paragraph()
    p_sig.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p_sig.add_run("Authorized Signatory\nDirector of Talent Acquisition\nFillForge Technologies")

    doc.save(str(dest_path))


def seed_samples():
    init_storage()
    init_db()

    existing = get_all_documents()
    existing_names = {d["name"] for d in existing}

    sample_files = [
        ("Offer Letter.docx", create_offer_letter_template),
        ("Internship Certificate.docx", create_internship_certificate_template),
    ]

    for name, creator in sample_files:
        if name not in existing_names:
            doc_id = str(uuid.uuid4())
            filename = f"{doc_id}_{name}"
            target_path = get_template_path(filename)
            creator(target_path)
            placeholders = extract_placeholders_from_docx(target_path)
            insert_document(doc_id, name, filename, placeholders)
            print(f"Seeded sample template: {name} (ID: {doc_id})")

if __name__ == "__main__":
    seed_samples()
